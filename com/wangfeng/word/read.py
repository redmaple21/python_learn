import docx

file = docx.Document('b.docx')
result = []
for table in file.tables:
    for index in range(1, len(table.rows)):
        row = table.rows[index]
        value = []
        for cell in row.cells:
            if cell.text not in value:
                value.append(cell.text)
        print(value)
        result.append(value)
print(result)
print('read done!')

import xlwt

work_book = xlwt.Workbook(encoding='urf-8')
sheet = work_book.add_sheet('数据统计')
title = ['姓名', '年龄', '性别', '身高']
for index in range(0, len(title)):
    sheet.write(0, index, title[index])
for i in range(0, len(result)):
    for j in range(0, len(result[i])):
        sheet.write(i + 1, j, result[i][j])
work_book.save('b.xls')
print('write done!')
